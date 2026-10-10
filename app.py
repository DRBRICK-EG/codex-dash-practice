"""Local WooCommerce → GNT converter. Customer files are processed in memory."""
import base64
import io
import json
import re
from decimal import Decimal, InvalidOperation
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl
import xlwt

ROOT = Path(__file__).resolve().parent
SCHEMA = json.loads((ROOT / 'gnt_schema.json').read_text())
HEADERS = SCHEMA['headers']
LOOKUP = SCHEMA['lookup']
CITIES = {name: [row[i] for row in LOOKUP[1:] if row[i]] for i, name in enumerate(LOOKUP[0])}
# Only mappings verified for the supplied export. Other codes require manual selection.
STATES = {'EGGZ': 'الجيزة', 'EGC': 'القاهرة'}


def sender_settings():
    path = ROOT / 'sender-settings.json'
    return json.loads(path.read_text()) if path.exists() else {}


def string(value):
    if value is None:
        return ''
    return str(value).strip()


def norm(value):
    return re.sub(r'[\s\u064b-\u065f\u0640]+', '', value.translate(str.maketrans('أإآىة', 'ااايه')))


def parse_workbook(data):
    book = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    try:
        sheet = book.active
        if sheet.max_row > 5001 or sheet.max_column > 300:
            raise ValueError('每次最多处理 5000 行、300 列。')
        iterator = sheet.iter_rows(values_only=True)
        columns = [string(x) for x in next(iterator)]
        needed = {'order_number', 'status', 'payment_method', 'order_total'}
        if not needed.issubset(columns):
            raise ValueError('缺少 WooCommerce 列：' + ', '.join(sorted(needed - set(columns))))
        rows, warnings, seen = [], [], set()
        for source in iterator:
            if not any(v is not None for v in source):
                continue
            if any(isinstance(v, str) and v.startswith('=') for v in source):
                raise ValueError('文件包含公式，请导出仅包含数值的订单文件。')
            d = dict(zip(columns, source))
            get = lambda key: string(d.get(key))
            number = get('order_number')
            if not number or number in seen:
                raise ValueError('订单编号为空或重复；请使用每笔订单一行的导出格式。')
            seen.add(number)
            if get('status') not in ('processing', 'wc-processing'):
                warnings.append(f'订单 {number}：状态不是 processing，已跳过。')
                continue
            if get('refund_items'):
                raise ValueError(f'订单 {number} 有退款记录，需要先人工确认应收金额。')
            row = [''] * len(HEADERS)
            row[0] = number
            shipping = bool(get('shipping_address_1'))
            prefix = 'shipping_' if shipping else 'billing_'
            row[1] = ' '.join(filter(None, [get(prefix + 'first_name'), get(prefix + 'last_name')]))
            row[2] = get('shipping_phone') or get('billing_phone')
            if not isinstance(d.get('shipping_phone') or d.get('billing_phone'), str):
                warnings.append(f'订单 {number}：电话在源文件中是数字，已转为文本；请核对国家码和前导 0，工具不会猜测补号。')
            state = get(prefix + 'state')
            row[4] = STATES.get(state, state)
            city = get(prefix + 'city')
            matches = [c for c in CITIES.get(row[4], []) if norm(c) == norm(city)]
            row[5] = matches[0] if len(matches) == 1 else city
            if row[5] != city:
                warnings.append(f'订单 {number}：城市名称已从 {city} 对应为 {row[5]}，请核对。')
            row[7] = ' / '.join(filter(None, [get(prefix + 'address_1'), get(prefix + 'address_2')]))
            if get(prefix + 'country') != 'EG':
                raise ValueError(f'订单 {number} 不是埃及地址，不能使用当前模板。')
            items = []
            for key in columns:
                if re.fullmatch(r'line_item_\d+', key) and get(key):
                    parts = dict(piece.split(':', 1) for piece in get(key).split('|') if ':' in piece)
                    if not parts.get('name') or not parts.get('quantity'):
                        raise ValueError(f'订单 {number} 商品格式无法识别。')
                    items.append(parts['name'] + ' × ' + parts['quantity'])
            row[16] = '; '.join(items)
            row[20] = get('order_currency')
            if get('payment_method') == 'cod' and not get('paid_date'):
                row[21] = get('order_total')
            elif get('payment_method') != 'cod' and get('paid_date'):
                row[21] = '0'
                warnings.append(f'订单 {number}：有付款日期，代收金额设为 0；请确认到账。')
            else:
                warnings.append(f'订单 {number}：付款状态需要人工确认，请填写正确的 COD amount。')
            row[26] = get('customer_note')
            for index, value in sender_settings().items():
                if int(index) in range(8, 15):
                    row[int(index)] = value
            rows.append(row)
        if not rows:
            raise ValueError('没有可转换的 processing 订单。')
        return {'headers': HEADERS, 'rows': rows, 'warnings': warnings, 'cities': CITIES, 'sender': sender_settings()}
    finally:
        book.close()


def validate(rows):
    errors, seen = [], set()
    if not rows:
        return ['没有订单。']
    for index, row in enumerate(rows, 1):
        if len(row) != len(HEADERS):
            errors.append(f'第 {index} 行列数错误。')
            continue
        label = f'第 {index} 行（订单 {row[0]}）'
        if not row[0] or row[0] in seen:
            errors.append(label + '：订单编号为空或重复。')
        seen.add(row[0])
        for i, header in enumerate(HEADERS):
            if header.startswith('*') and not string(row[i]):
                errors.append(label + '：缺少 ' + header)
        for a, b in [(4, 5), (11, 12)]:
            if row[a] not in CITIES or row[b] not in CITIES.get(row[a], []):
                errors.append(label + '：省／市不符合模板，请核对 ' + HEADERS[a])
        if row[20] != 'EGP':
            errors.append(label + '：当前版本只支持 EGP。')
        for i in (17, 21, 22, 23):
            if i != 21 and not string(row[i]):
                continue
            try:
                amount = Decimal(string(row[i]))
                if not amount.is_finite() or amount < 0:
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                errors.append(label + '：金额／重量无效 ' + HEADERS[i])
    return errors


def export_xls(rows):
    errors = validate(rows)
    if errors:
        raise ValueError('\n'.join(errors))
    book = xlwt.Workbook()
    sheet = book.add_sheet('Sheet1')
    for i, header in enumerate(HEADERS):
        sheet.write(0, i, header)
        sheet.col(i).width = 5500
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            # Strings stay strings, including phone numbers and formula-like customer input.
            if c in (17, 21, 22, 23) and string(value):
                sheet.write(r, c, float(Decimal(string(value))))
            else:
                sheet.write(r, c, string(value))
    lookup = book.add_sheet('Sheet2')
    for r, row in enumerate(LOOKUP):
        for c, value in enumerate(row):
            lookup.write(r, c, value)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def reply(self, status, data, kind='application/json; charset=utf-8'):
        self.send_response(status)
        self.send_header('Content-Type', kind)
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        if self.path == '/':
            self.reply(200, (ROOT / 'index.html').read_bytes(), 'text/html; charset=utf-8')
        else:
            self.reply(404, b'{}')

    def do_POST(self):
        try:
            if self.headers.get('Origin') and self.headers['Origin'] != 'http://' + self.headers.get('Host', ''):
                raise ValueError('不允许跨站请求。')
            size = int(self.headers.get('Content-Length', 0))
            if not 0 < size <= 12 * 1024 * 1024:
                raise ValueError('文件或请求过大，限 12 MB。')
            payload = json.loads(self.rfile.read(size))
            if self.path == '/parse':
                result = parse_workbook(base64.b64decode(payload['file'], validate=True))
                self.reply(200, json.dumps(result, ensure_ascii=False).encode())
            elif self.path == '/export':
                if payload.get('confirmed') is not True:
                    raise ValueError('请先核对区域、金额和服务字段。')
                self.reply(200, export_xls(payload['rows']), 'application/vnd.ms-excel')
            else:
                self.reply(404, b'{}')
        except Exception as e:
            message = str(e) if isinstance(e, ValueError) else '文件无法处理，请检查格式和内容。'
            self.reply(400, json.dumps({'error': message}, ensure_ascii=False).encode())


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--port', type=int, default=8765)
    parser.add_argument('--open-browser', action='store_true')
    args = parser.parse_args()
    server = ThreadingHTTPServer(('127.0.0.1', args.port), Handler)
    if args.open_browser:
        import webbrowser
        webbrowser.open(f'http://127.0.0.1:{args.port}')
    print(f'Local converter started on port {args.port}', flush=True)
    server.serve_forever()
